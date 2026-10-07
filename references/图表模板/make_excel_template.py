"""从同源 CSV 生成可编辑的 Excel XY 图表，无外部链接或公式。"""
import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import Reference, ScatterChart, Series
from openpyxl.styles import Alignment, Font, PatternFill


def build(root):
    with (root / "a175_mesh.csv").open(encoding="utf-8-sig", newline="") as handle:
        rows = sorted((float(r["mesh_step_m"]), float(r["power_kw_m2"]))
                      for r in csv.DictReader(handle))
    book = Workbook()
    data = book.active
    data.title = "数据"
    data.append(["网格步长 (m)", "单位镜面面积功率 (kW/m²)"])
    for row in rows:
        data.append(row)
    data.freeze_panes = "A2"
    data.column_dimensions["A"].width = 22
    data.column_dimensions["B"].width = 34
    for cell in data[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="276FBF")
    for row in data.iter_rows(min_row=2):
        row[0].number_format = "0.##"
        row[1].number_format = "0.00000000"
    chart = ScatterChart()
    chart.scatterStyle = "lineMarker"
    chart.title = "网格分辨率敏感性"
    chart.x_axis.title = "网格步长 (m)"
    chart.y_axis.title = "单位镜面面积功率 (kW/m²)"
    series = Series(Reference(data, min_col=2, min_row=2, max_row=len(rows) + 1),
                    Reference(data, min_col=1, min_row=2, max_row=len(rows) + 1),
                    title="A175 表3")
    series.graphicalProperties.line.solidFill = "276FBF"
    series.graphicalProperties.line.width = 19050
    series.marker.symbol = "circle"
    series.marker.size = 5
    series.marker.graphicalProperties.solidFill = "276FBF"
    series.marker.graphicalProperties.line.solidFill = "276FBF"
    series.smooth = False
    chart.series.append(series)
    chart.legend = None
    chart.width, chart.height = 17, 11
    data.add_chart(chart, "D2")
    notes = book.create_sheet("说明")
    for line in ["来源：2023 A175，PDF p.14 表3。仅复画论文表内数值，未重新求解赛题。",
                 "XY 散点连线保留不等距步长；不能换成等间距分类折线图。",
                 "替换数据时同步修改轴名、单位、来源和图题；增加行数后编辑图表数据范围。",
                 "修改工作表数值后，Excel 图表引用随之更新。没有公式、宏或外部链接。",
                 "0.25 m 的结果只是本表最细网格参考，不能称为真值或精确解。",
                 "图表结构与引用经 openpyxl 校验；未在桌面 Excel 中验证渲染。"]:
        notes.append([line])
    notes.column_dimensions["A"].width = 105
    for row in notes:
        row[0].alignment = Alignment(wrap_text=True, vertical="top")
        notes.row_dimensions[row[0].row].height = 32
    book.save(root / "excel_template.xlsx")


if __name__ == "__main__":
    build(Path(__file__).resolve().parent)
